import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def export_to_excel(header, items, output_path):

    wb = Workbook()

    # ==========================================
    # HEADER SHEET
    # ==========================================

    ws1 = wb.active
    ws1.title = "Header"

    header_fill = PatternFill(
        start_color="4F81BD",
        end_color="4F81BD",
        fill_type="solid"
    )

    bold = Font(bold=True)

    ws1["A1"] = "Field"
    ws1["B1"] = "Value"

    ws1["A1"].font = bold
    ws1["B1"].font = bold

    ws1["A1"].fill = header_fill
    ws1["B1"].fill = header_fill

    row = 2

    for key, value in header.items():

        ws1.cell(row=row, column=1).value = key
        ws1.cell(row=row, column=2).value = value

        row += 1

    ws1.column_dimensions["A"].width = 30
    ws1.column_dimensions["B"].width = 45

    # ==========================================
    # ITEM SHEET
    # ==========================================

    ws2 = wb.create_sheet("Items")

    columns = [

        "Description",
        "HSN",
        "Quantity",
        "Unit",
        "Rate",
        "Amount"

    ]

    for col, name in enumerate(columns, start=1):

        cell = ws2.cell(row=1, column=col)

        cell.value = name
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    row = 2

    for item in items:

        ws2.cell(row=row, column=1).value = item.get("Description", "")
        ws2.cell(row=row, column=2).value = item.get("HSN", "")
        ws2.cell(row=row, column=3).value = item.get("Quantity", "")
        ws2.cell(row=row, column=4).value = item.get("Per", "")
        ws2.cell(row=row, column=5).value = item.get("Rate", "")
        ws2.cell(row=row, column=6).value = item.get("Amount", "")

        row += 1

    widths = {
        "A": 45,
        "B": 15,
        "C": 15,
        "D": 15,
        "E": 15,
        "F": 18
    }

    for col, width in widths.items():

        ws2.column_dimensions[col].width = width

    # ==========================================
    # SAVE
    # ==========================================

    output_dir = os.path.dirname(output_path)

    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    wb.save(output_path)

    print("\n==========================================")
    print("Excel exported successfully")
    print(output_path)
    print("==========================================")